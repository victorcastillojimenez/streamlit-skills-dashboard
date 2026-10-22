import numpy as np
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

from utils.data_loader import generate_mock_data

def load_data():
    df = generate_mock_data()
    df["Fecha"] = df["Fecha"].apply(lambda x: x.strftime("%Y-%m-%d"))
    return df.to_dict('records')

def generate_excel():
    data = load_data()
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos Dashboard"
    
    # Encabezados
    headers = ["Fecha", "Categoría", "Ventas", "Visitas"]
    ws.append(headers)
    
    # Estilos de encabezado
    header_font = Font(name='Arial', bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='004C99')
    for col_num in range(1, 5):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    # Datos
    for row_idx, row_data in enumerate(data, start=2):
        ws.cell(row=row_idx, column=1, value=row_data["Fecha"])
        ws.cell(row=row_idx, column=2, value=row_data["Categoría"])
        ws.cell(row=row_idx, column=3, value=row_data["Ventas"])
        ws.cell(row=row_idx, column=4, value=row_data["Visitas"])
        
    last_row = len(data) + 1
    
    # Fila de Totales con fórmulas Excel
    ws.cell(row=last_row + 2, column=2, value="TOTALES").font = Font(name='Arial', bold=True)
    
    # Fórmulas
    sum_ventas_cell = ws.cell(row=last_row + 2, column=3, value=f"=SUM(C2:C{last_row})")
    sum_ventas_cell.font = Font(name='Arial', bold=True)
    
    sum_visitas_cell = ws.cell(row=last_row + 2, column=4, value=f"=SUM(D2:D{last_row})")
    sum_visitas_cell.font = Font(name='Arial', bold=True)
    
    # Formato de números
    for row in range(2, last_row + 3):
        ws.cell(row=row, column=3).number_format = '"$"#,##0'
        ws.cell(row=row, column=4).number_format = '#,##0'
        
        # Fuente general
        for col in range(1, 5):
            ws.cell(row=row, column=col).font = Font(name='Arial')

    # Ancho de columnas
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    
    output_path = "dashboard_datos.xlsx"
    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    generate_excel()
